"""Generación de archivos Excel con formato profesional y respaldos completos."""

from __future__ import annotations

import io
import json
import zipfile
from datetime import datetime

import pandas as pd
import sqlalchemy as sa
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import database as db
import reports
from config import now_lima

MONEY = '"S/" #,##0.00'
DATE = "DD/MM/YYYY"
DATETIME = "DD/MM/YYYY HH:MM"
TIME = "hh:mm AM/PM"
INTEGER = "0"

HEADER_FILL = PatternFill("solid", fgColor="1F4E5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FILL = PatternFill("solid", fgColor="E8F0F2")
PENDING_FONT = Font(color="B42318")
PAID_FONT = Font(color="067647")
THIN = Side(style="thin", color="D0D7DE")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _write_sheet(wb: Workbook, title: str, df: pd.DataFrame, formats: dict[str, str], totals: list[str] | None = None):
    ws = wb.create_sheet(title)
    columns = list(df.columns)
    ws.append(columns)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = 30

    for record in df.itertuples(index=False):
        ws.append([_excel_value(v) for v in record])

    n_rows = len(df)
    status_col = columns.index("Estado") + 1 if "Estado" in columns else None
    for row in ws.iter_rows(min_row=2, max_row=n_rows + 1):
        for cell in row:
            cell.border = BORDER
            fmt = formats.get(columns[cell.column - 1])
            if fmt:
                cell.number_format = fmt
            if cell.column == status_col:
                cell.font = PAID_FONT if cell.value == "Pagado" else PENDING_FONT

    last_col = get_column_letter(len(columns))
    ws.auto_filter.ref = f"A1:{last_col}{max(n_rows + 1, 2)}"
    ws.freeze_panes = "A2"

    if totals and n_rows:
        total_row = n_rows + 2
        ws.cell(row=total_row, column=1, value="TOTAL")
        for col_name in totals:
            idx = columns.index(col_name) + 1
            letter = get_column_letter(idx)
            cell = ws.cell(row=total_row, column=idx, value=f"=SUBTOTAL(9,{letter}2:{letter}{n_rows + 1})")
            cell.number_format = formats.get(col_name, MONEY)
        for cell in ws[total_row]:
            cell.fill = TOTAL_FILL
            cell.font = Font(bold=True)
            cell.border = BORDER

    for idx, col_name in enumerate(columns, start=1):
        values = [len(str(col_name))] + [len(_display(v)) for v in df.iloc[:, idx - 1].head(500)]
        ws.column_dimensions[get_column_letter(idx)].width = min(max(values) + 3, 50)
    return ws


def _excel_value(value):
    if value is None or (not isinstance(value, str) and pd.isna(value)):
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if isinstance(value, pd.Period):
        return reports.month_label(value)
    return value


def _display(value) -> str:
    if value is None or (not isinstance(value, str) and pd.isna(value)):
        return ""
    if isinstance(value, (pd.Timestamp, datetime)):
        return "00/00/0000 00:00"
    return str(value)


def _general_sheet(records: pd.DataFrame) -> pd.DataFrame:
    return pd.DataFrame({
        "Fecha de actividad": records["activity_date"],
        "Fecha de registro": records["registered_at"],
        "Actividad": records["activity"],
        "Voluntario": records["volunteer"],
        "Hora de inicio": records["start_time"],
        "Hora de llegada": records["arrival_time"],
        "Minutos de retraso": records["delay_minutes"],
        "Penalidad": records["penalty"],
        "Estado": records["status"],
        "Monto pagado": records["paid"],
        "Saldo pendiente": records["balance"],
        "Fecha de pago": records["payment_date"],
        "Medio de pago": records["payment_method"],
        "Observaciones": [
            " | ".join(x for x in (n, p) if isinstance(x, str) and x)
            for n, p in zip(records["notes"], records["payment_notes"])
        ],
    })


def build_workbook(records: pd.DataFrame, pays: pd.DataFrame) -> bytes:
    """Crea el Excel de 5 hojas a partir de los registros (completos o filtrados)."""
    records = records.sort_values(["activity_date", "activity", "arrival_time"])
    wb = Workbook()
    wb.remove(wb.active)

    _write_sheet(
        wb, "Registro general", _general_sheet(records),
        {"Fecha de actividad": DATE, "Fecha de registro": DATETIME, "Hora de inicio": TIME, "Hora de llegada": TIME,
         "Minutos de retraso": INTEGER, "Penalidad": MONEY, "Monto pagado": MONEY, "Saldo pendiente": MONEY,
         "Fecha de pago": DATE},
        totals=["Penalidad", "Monto pagado", "Saldo pendiente"],
    )

    by_activity = reports.activity_summary(records).drop(columns=["activity_id"]).sort_values(["Fecha", "Actividad"])
    by_activity = by_activity.rename(columns={"Penalidades": "Total penalidades"})
    _write_sheet(
        wb, "Resumen por actividad", by_activity,
        {"Fecha": DATE, "Hora de inicio": TIME, "Total penalidades": MONEY, "Recaudado": MONEY, "Pendiente": MONEY},
        totals=["Tardanzas", "Total penalidades", "Recaudado", "Pendiente"],
    )

    monthly = reports.monthly_summary(records).drop(columns=["month"])
    monthly = monthly.rename(columns={"Penalidades": "Penalidades generadas", "Recaudado": "Monto recaudado",
                                      "Pendiente": "Monto pendiente"})
    _write_sheet(
        wb, "Resumen mensual", monthly,
        {"Penalidades generadas": MONEY, "Monto recaudado": MONEY, "Monto pendiente": MONEY},
        totals=["Tardanzas", "Penalidades generadas", "Monto recaudado", "Monto pendiente"],
    )

    pending = records[records["status"] == db.STATUS_PENDING].sort_values(["volunteer", "activity_date"])
    _write_sheet(
        wb, "Pagos pendientes",
        pd.DataFrame({
            "Voluntario": pending["volunteer"],
            "Fecha de actividad": pending["activity_date"],
            "Actividad": pending["activity"],
            "Penalidad": pending["penalty"],
            "Pagado a cuenta": pending["paid"],
            "Monto pendiente": pending["balance"],
        }),
        {"Fecha de actividad": DATE, "Penalidad": MONEY, "Pagado a cuenta": MONEY, "Monto pendiente": MONEY},
        totals=["Penalidad", "Pagado a cuenta", "Monto pendiente"],
    )

    pays = pays[pays["tardy_id"].isin(records["id"])].sort_values("payment_date")
    _write_sheet(
        wb, "Historial de pagos",
        pd.DataFrame({
            "Fecha de pago": pays["payment_date"],
            "Voluntario": pays["volunteer"],
            "Actividad": pays["activity"],
            "Fecha de actividad": pays["activity_date"],
            "Monto": pays["amount"],
            "Medio de pago": pays["method"],
            "Observaciones": pays["notes"],
            "Registrado el": pays["created_at"],
        }),
        {"Fecha de pago": DATE, "Fecha de actividad": DATE, "Monto": MONEY, "Registrado el": DATETIME},
        totals=["Monto"],
    )

    wb.properties.title = "Reporte de tardanzas"
    wb.properties.creator = "Control de Tardanzas"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def file_name(suffix: str = "") -> str:
    stamp = now_lima().strftime("%Y-%m-%d_%H%M")
    suffix = f"_{suffix}" if suffix else ""
    return f"tardanzas{suffix}_{stamp}.xlsx"


def backup_zip() -> bytes:
    """Copia de seguridad completa: cada tabla en CSV y JSON, incluido el historial de cambios."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf, db.get_engine().connect() as conn:
        manifest = {"generado": now_lima().isoformat(), "tablas": {}}
        for table in db.metadata.sorted_tables:
            rows = [dict(r) for r in conn.execute(sa.select(table)).mappings().all()]
            df = pd.DataFrame(rows, columns=[c.name for c in table.columns])
            zf.writestr(f"{table.name}.csv", df.to_csv(index=False))
            zf.writestr(f"{table.name}.json", json.dumps(rows, ensure_ascii=False, default=str, indent=1))
            manifest["tablas"][table.name] = len(rows)
        zf.writestr("manifest.json", json.dumps(manifest, ensure_ascii=False, indent=1))
    return buffer.getvalue()
