"""Pruebas de cálculos y flujos principales con una base SQLite temporal.

Ejecutar con:  python -m pytest -q
"""

import io
import os
import sys
from datetime import date, time, timedelta
from decimal import Decimal
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))


@pytest.fixture()
def app(tmp_path, monkeypatch):
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{tmp_path / 'test.db'}")
    import config
    import excel_export
    import payments
    import reports
    import tardiness
    import volunteers

    return config, tardiness, payments, volunteers, reports, excel_export


# ------------------------------------------------------------------ cálculo


@pytest.mark.parametrize(
    "arrival, minutes, penalty",
    [
        (time(8, 55), 0, 0), (time(9, 0), 0, 0), (time(9, 5), 5, 0), (time(9, 9), 9, 0),
        (time(9, 10), 10, 1), (time(9, 15), 15, 1), (time(9, 19), 19, 1),
        (time(9, 20), 20, 2), (time(9, 25), 25, 2), (time(9, 29), 29, 2),
        (time(9, 30), 30, 3), (time(9, 45), 45, 3), (time(11, 0), 120, 3),
    ],
)
def test_penalty_intervals(arrival, minutes, penalty):
    from config import DEFAULT_PENALTY_RULES, delay_minutes, penalty_for_delay

    assert delay_minutes(time(9, 0), arrival) == minutes
    assert penalty_for_delay(minutes, DEFAULT_PENALTY_RULES) == Decimal(penalty)


def test_rules_validation():
    from config import validate_rules

    with pytest.raises(ValueError):
        validate_rules([{"desde_min": 5, "monto": 1}])
    with pytest.raises(ValueError):
        validate_rules([{"desde_min": 0, "monto": 0}, {"desde_min": 0, "monto": 1}])
    assert validate_rules([{"desde_min": 15, "monto": 2}, {"desde_min": 0, "monto": 0}])[0]["desde_min"] == 0


def test_describe_rules():
    from config import DEFAULT_PENALTY_RULES, describe_rules

    lines = describe_rules(DEFAULT_PENALTY_RULES, time(9, 0))
    assert lines[0] == "Antes de las 9:10 a. m. → S/ 0.00"
    assert lines[-1] == "Desde las 9:30 a. m. en adelante → S/ 3.00"


# ------------------------------------------------------------------- flujos


def test_batch_registration_and_payments(app):
    config, tardiness, payments, volunteers, reports, excel_export = app
    yesterday = date(2026, 10, 29)
    act_id, created, start = tardiness.get_or_create_activity("Reunión general", yesterday, time(9, 0))
    assert created and start == time(9, 0)

    # Registro consecutivo en la misma actividad
    r1 = tardiness.register_tardy(act_id, "Ana", time(9, 12), client_token="t1")
    r2 = tardiness.register_tardy(act_id, "Pedro", time(9, 25), client_token="t2")
    r3 = tardiness.register_tardy(act_id, "María", time(9, 35), client_token="t3")
    assert (r1["penalty"], r2["penalty"], r3["penalty"]) == (1, 2, 3)
    assert (r1["delay_minutes"], r2["delay_minutes"], r3["delay_minutes"]) == (12, 25, 35)

    # Reintento del mismo envío: no duplica
    again = tardiness.register_tardy(act_id, "Ana", time(9, 12), client_token="t1")
    assert again["duplicate"] and again["id"] == r1["id"]
    # Mismo voluntario (con otra escritura) en la misma actividad: se rechaza
    with pytest.raises(tardiness.TardyError):
        tardiness.register_tardy(act_id, "  ana ", time(9, 40), client_token="t4")
    # Llegada puntual: no es tardanza
    with pytest.raises(tardiness.TardyError):
        tardiness.register_tardy(act_id, "Luis", time(9, 0), client_token="t5")

    # Misma actividad + fecha reutiliza la actividad
    same_id, created, _ = tardiness.get_or_create_activity("reunion general", yesterday, time(10, 0))
    assert same_id == act_id and not created

    df = tardiness.fetch_records()
    assert len(df) == 3 and set(df["status"]) == {"Pendiente"}
    assert (df["registered_at"].dt.date == config.today_lima()).all()  # fecha de registro ≠ fecha de actividad
    assert (df["activity_date"].dt.date == yesterday).all()

    # Pago parcial y luego total
    status = payments.register_payment(r3["id"], config.today_lima(), 1, "Yape", "a cuenta")
    assert status == "Pendiente"
    with pytest.raises(payments.PaymentError):
        payments.register_payment(r3["id"], config.today_lima(), 5, "Yape")
    status = payments.register_payment(r3["id"], config.today_lima(), 2, "Efectivo")
    assert status == "Pagado"
    payments.register_payments([r1["id"]], config.today_lima(), "Plin")

    df = tardiness.fetch_records()
    k = reports.kpis(df)
    assert k["generated"] == 6 and k["collected"] == 4 and k["pending"] == 2
    assert round(k["collection_rate"], 2) == 66.67
    maria = df[df["volunteer"] == "María"].iloc[0]
    assert maria["payment_method"] == "Yape, Efectivo"

    # Anular un pago devuelve la deuda a pendiente
    pays = payments.fetch_payments()
    plin = pays[pays["method"] == "Plin"].iloc[0]
    payments.void_payment(int(plin["payment_id"]), "registrado por error")
    df = tardiness.fetch_records()
    assert df.set_index("id").loc[r1["id"], "status"] == "Pendiente"
    assert reports.kpis(df)["pending"] == 3

    # No se puede anular una tardanza con pagos
    with pytest.raises(tardiness.TardyError):
        tardiness.void_tardy(r3["id"], "error")
    tardiness.void_tardy(r2["id"], "no llegó tarde")
    assert len(tardiness.fetch_records()) == 2
    # Tras anular, se puede volver a registrar a esa persona
    tardiness.register_tardy(act_id, "Pedro", time(9, 21), client_token="t6")


def test_rule_changes_do_not_touch_history(app):
    config, tardiness, payments, volunteers, reports, excel_export = app
    act_id, *_ = tardiness.get_or_create_activity("Taller", date(2026, 9, 1), time(9, 0))
    old = tardiness.register_tardy(act_id, "Ana", time(9, 15), client_token="a")
    assert old["penalty"] == 1

    config.save_settings({"penalty_rules": [{"desde_min": 0, "monto": 0}, {"desde_min": 5, "monto": 5}],
                          "default_start_time": "08:30"})
    new = tardiness.register_tardy(act_id, "Pedro", time(9, 15), client_token="b")
    assert new["penalty"] == 5
    df = tardiness.fetch_records().set_index("id")
    assert df.loc[old["id"], "penalty"] == 1  # el histórico conserva su penalidad

    # Editar la llegada usa las reglas originales del registro
    tardiness.update_tardy(old["id"], time(9, 22), None, "hora mal digitada")
    assert tardiness.fetch_records().set_index("id").loc[old["id"], "penalty"] == 2

    # Recalcular solo a pedido expreso
    rules = config.get_settings()["penalty_rules"]
    preview = tardiness.preview_recalculation([old["id"]], rules)
    assert preview.iloc[0]["penalty_after"] == 5
    assert tardiness.recalculate_with_rules([old["id"]], rules) == 1
    assert tardiness.fetch_records().set_index("id").loc[old["id"], "penalty"] == 5

    # Corregir la hora de inicio de la actividad recalcula el retraso
    tardiness.update_activity(act_id, "Taller", date(2026, 9, 1), time(9, 10))
    df = tardiness.fetch_records().set_index("id")
    assert df.loc[new["id"], "delay_minutes"] == 5


def test_volunteers_and_filters(app):
    config, tardiness, payments, volunteers, reports, excel_export = app
    vid, created = volunteers.create_volunteer("José Pérez")
    assert created
    assert volunteers.create_volunteer("jose  perez") == (vid, False)
    with pytest.raises(volunteers.VolunteerError):
        volunteers.create_volunteer("x")
    volunteers.rename_volunteer(vid, "José Pérez Ruiz")

    a1, *_ = tardiness.get_or_create_activity("Reunión", date(2026, 8, 10), time(9, 0))
    a2, *_ = tardiness.get_or_create_activity("Campaña", date(2026, 10, 5), time(8, 0))
    tardiness.register_tardy(a1, "José Pérez Ruiz", time(9, 31))
    tardiness.register_tardy(a2, "José Pérez Ruiz", time(8, 11))
    tardiness.register_tardy(a2, "Ana", time(8, 20))
    with pytest.raises(volunteers.VolunteerError):
        volunteers.delete_volunteer(vid)

    df = tardiness.fetch_records()
    assert len(reports.filter_records(df, volunteers=["José Pérez Ruiz"])) == 2
    assert len(reports.filter_records(df, activities=["Campaña"])) == 2
    assert len(reports.filter_records(df, date_from=date(2026, 10, 1), date_to=date(2026, 10, 31))) == 2
    assert len(reports.filter_records(df, months=[df["month"].min()])) == 1
    assert len(reports.filter_records(df, status=["Pagado"])) == 0

    summary = reports.volunteer_summary(df, volunteers.list_volunteers())
    jose = summary[summary["volunteer"] == "José Pérez Ruiz"].iloc[0]
    assert jose["Tardanzas"] == 2 and jose["Total"] == 4 and jose["Pendiente"] == 4
    monthly = reports.monthly_summary(df)
    assert list(monthly["Mes"]) == ["Agosto 2026", "Octubre 2026"]

    # Excel completo y filtrado
    from openpyxl import load_workbook

    content = excel_export.build_workbook(df, payments.fetch_payments())
    wb = load_workbook(io.BytesIO(content))
    assert wb.sheetnames == ["Registro general", "Resumen por actividad", "Resumen mensual",
                             "Pagos pendientes", "Historial de pagos"]
    ws = wb["Registro general"]
    assert ws.max_row == 1 + 3 + 1  # encabezado + registros + total
    assert ws["H2"].number_format == '"S/" #,##0.00'
    assert ws.auto_filter.ref.startswith("A1:")
    filtered = excel_export.build_workbook(reports.filter_records(df, activities=["Campaña"]), payments.fetch_payments())
    assert load_workbook(io.BytesIO(filtered))["Registro general"].max_row == 1 + 2 + 1

    backup = excel_export.backup_zip()
    import zipfile

    assert "tardies.csv" in zipfile.ZipFile(io.BytesIO(backup)).namelist()


def test_batch_save_is_all_or_nothing(app):
    config, tardiness, payments, volunteers, reports, excel_export = app
    act_id, *_ = tardiness.get_or_create_activity("Reunión", date(2026, 10, 1), time(9, 0))
    entries = [{"name": "Ana", "arrival": time(9, 12)}, {"name": "Pedro", "arrival": time(9, 25)},
               {"name": "María", "arrival": time(9, 35)}]
    results = tardiness.register_tardies(act_id, entries, batch_token="b1")
    assert [r["penalty"] for r in results] == [1, 2, 3]
    # Reintento del mismo lote (p. ej. tras un corte de conexión): no duplica
    assert all(r["duplicate"] for r in tardiness.register_tardies(act_id, entries, batch_token="b1"))
    assert len(tardiness.fetch_records()) == 3

    # Si un registro del lote falla, no se guarda ninguno
    bad = [{"name": "Luis", "arrival": time(9, 15)}, {"name": "Ana", "arrival": time(9, 40)}]
    with pytest.raises(tardiness.TardyError, match="Ana"):
        tardiness.register_tardies(act_id, bad, batch_token="b2")
    assert "Luis" not in set(tardiness.fetch_records()["volunteer"])
    with pytest.raises(tardiness.TardyError):
        tardiness.register_tardies(act_id, [{"name": "Rosa", "arrival": time(9, 15)},
                                            {"name": "rosa", "arrival": time(9, 20)}])

    # Consulta por fecha y PDF
    import pdf_export

    df = tardiness.fetch_records()
    day = df[df["activity_date"].dt.date == date(2026, 10, 1)]
    assert list(day.sort_values("arrival_time")["volunteer"]) == ["Ana", "Pedro", "María"]
    content = pdf_export.build_pdf(day, "Reporte de tardanzas", "Fecha de actividad: 01/10/2026", "Voluntariado")
    assert content.startswith(b"%PDF") and len(content) > 1000


def test_future_dates_rejected(app):
    config, tardiness, *_ = app
    with pytest.raises(tardiness.TardyError):
        tardiness.get_or_create_activity("Futuro", config.today_lima() + timedelta(days=1), time(9, 0))
